#!/usr/bin/env python3
"""UNOVA AI: inventory, classify, extract text, and flag files for review.

Usage (from project root):
  python organize_unova.py --input data/raw --output data/processed --report data/sources.csv

Dependencies:
  python -m pip install pymupdf python-docx

This script does NOT decide whether a document is legally current. It creates a
first-pass classification and flags uncertain, unreadable, duplicate, or possibly
sensitive files for human review.
"""
from __future__ import annotations
import argparse, csv, hashlib, re, sys
from pathlib import Path
from datetime import date

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None
try:
    from docx import Document
except ImportError:
    Document = None

SUPPORTED = {'.pdf', '.docx', '.txt', '.doc', '.png', '.jpg', '.jpeg', '.webp'}

CATEGORY_RULES = [
    ('akademik_takvim', ['akademik takvim', 'akademik_takvim', 'egitim ogretim takvimi', 'eğitim öğretim takvimi']),
    ('ogrenci_isleri', ['ogrenci isleri', 'öğrenci işleri', 'ders kayit', 'ders kayıt', 'kayit yenileme', 'kayıt yenileme', 'yatay gecis', 'yatay geçiş', 'muafiyet', 'intibak', 'mazeret sinav', 'mazeret sınav', 'diploma', 'mezuniyet', 'ogrenci belgesi', 'öğrenci belgesi']),
    ('mevzuat', ['yonetmelik', 'yönetmelik', 'yonerge', 'yönerge', 'usul ve esas', 'senato karari', 'senato kararı', 'not donusum', 'not dönüşüm', 'cift anadal', 'çift anadal', 'yan dal', 'yandal']),
    ('erasmus_uluslararasi', ['erasmus', 'international', 'uluslararasi', 'uluslararası', 'learning agreement', 'ogrenim hareketliligi', 'öğrenim hareketliliği']),
    ('kampus_yasami', ['kampus', 'kampüs', 'yemekhane', 'yemek bursu', 'kutuphane', 'kütüphane', 'spor tesisi', 'hali saha', 'halı saha', 'misafirhane', 'topluluk', 'ulasim', 'ulaşım', 'sks', 'saglikli yasam', 'sağlıklı yaşam']),
    ('personel_islemleri', ['ogretim elemani', 'öğretim elemanı', 'personel', 'not giris', 'not giriş', 'yoklama', 'sinav islemleri', 'sınav işlemleri', 'disiplin sorusturma', 'disiplin soruşturma', 'akademisyen']),
    ('basvuru_kilavuzlari', ['basvuru', 'başvuru', 'kilavuz', 'kılavuz', 'tercih', 'aday ogrenci', 'aday öğrenci']),
]
SENSITIVE_HINTS = ['ogrenci belgesi', 'öğrenci belgesi', 'kimlik no', 'tc kimlik', 't.c. kimlik', 'pasaport no', 'iban', 'telefon numarasi', 'telefon numarası', 'adres:']
STAFF_HINTS = ['personel', 'ogretim elemani', 'öğretim elemanı', 'disiplin sorusturma', 'disiplin soruşturma', 'not giris', 'not giriş', 'yoklama']


def norm(s: str) -> str:
    s = s.lower().replace('ı','i').replace('İ','i').replace('ş','s').replace('ğ','g').replace('ü','u').replace('ö','o').replace('ç','c')
    return re.sub(r'\s+', ' ', s)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open('rb') as f:
        for block in iter(lambda: f.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()


def extract(path: Path) -> tuple[str, str]:
    ext = path.suffix.lower()
    if ext == '.pdf':
        if fitz is None: return '', 'pymupdf_missing'
        try:
            with fitz.open(path) as pdf:
                pages = []
                for i, page in enumerate(pdf):
                    txt = page.get_text('text').strip()
                    if txt: pages.append(f'\n[Sayfa {i+1}]\n{txt}')
                text = '\n'.join(pages).strip()
                return text, ('ok' if len(text) >= 80 else 'ocr_or_manual_check')
        except Exception as e:
            return '', 'pdf_error:' + type(e).__name__
    if ext == '.docx':
        if Document is None: return '', 'python_docx_missing'
        try:
            doc = Document(path)
            chunks = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    chunks.append(' | '.join(c.text.strip() for c in row.cells))
            text = '\n'.join(chunks).strip()
            return text, ('ok' if text else 'empty_or_template_check')
        except Exception as e:
            return '', 'docx_error:' + type(e).__name__
    if ext == '.txt':
        for enc in ('utf-8', 'utf-8-sig', 'cp1254', 'latin-1'):
            try:
                text = path.read_text(encoding=enc).strip()
                return text, ('ok' if text else 'empty_text')
            except UnicodeDecodeError:
                continue
        return '', 'text_encoding_error'
    if ext == '.doc':
        return '', 'legacy_doc_convert_to_docx_or_pdf'
    if ext in {'.png', '.jpg', '.jpeg', '.webp'}:
        return '', 'image_ocr_or_manual_check'
    return '', 'unsupported'


def classify(path: Path, text: str) -> tuple[str, str, str]:
    combined = norm(str(path) + ' ' + text[:10000])
    for category, words in CATEGORY_RULES:
        if any(norm(w) in combined for w in words):
            confidence = 'medium' if any(norm(w) in norm(str(path)) for w in words) else 'low'
            audience = 'staff' if any(norm(w) in combined for w in STAFF_HINTS) else 'student'
            return category, confidence, audience
    return 'inceleme_gerekli', 'low', 'unknown'


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--input', default='data/raw', help='Folder containing extracted source files')
    ap.add_argument('--output', default='data/processed', help='Folder for extracted text')
    ap.add_argument('--report', default='data/sources.csv', help='CSV report path')
    args = ap.parse_args()
    root, out, report = Path(args.input), Path(args.output), Path(args.report)
    if not root.exists():
        print(f'Input folder not found: {root}\nCreate it and put the extracted ZIP contents inside.', file=sys.stderr)
        sys.exit(2)
    out.mkdir(parents=True, exist_ok=True)
    report.parent.mkdir(parents=True, exist_ok=True)
    files = [p for p in root.rglob('*') if p.is_file() and p.suffix.lower() in SUPPORTED
             and '__macosx' not in str(p).lower() and not p.name.startswith('._')]
    seen = {}
    rows = []
    for path in sorted(files):
        digest = sha256(path)
        text, extraction_status = extract(path)
        category, confidence, audience = classify(path, text)
        flags = []
        if digest in seen: flags.append('exact_duplicate_of:' + seen[digest])
        else: seen[digest] = str(path)
        ntext = norm(text[:20000])
        if any(norm(h) in norm(path.name + ' ' + text[:5000]) for h in SENSITIVE_HINTS):
            flags.append('possible_sensitive_data_review')
        if audience == 'staff': flags.append('staff_content_review_before_student_release')
        if extraction_status != 'ok': flags.append(extraction_status)
        if confidence == 'low' or category == 'inceleme_gerekli': flags.append('category_review')
        if not text and path.suffix.lower() not in {'.png','.jpg','.jpeg','.webp','.doc'}:
            flags.append('no_text_extracted')
        rel = path.relative_to(root)
        txt_rel = ''
        if text:
            safe_stem = re.sub(r'[^A-Za-z0-9._-]+', '_', rel.stem)[:100]
            target_dir = out / category
            target_dir.mkdir(parents=True, exist_ok=True)
            target = target_dir / f'{safe_stem}_{digest[:8]}.txt'
            target.write_text(f'SOURCE_FILE: {rel.as_posix()}\nCATEGORY: {category}\nSHA256: {digest}\n\n{text}\n', encoding='utf-8')
            txt_rel = str(target)
        rows.append({
            'document_id': digest[:12], 'filename': rel.as_posix(), 'file_type': path.suffix.lower(),
            'size_bytes': path.stat().st_size, 'sha256': digest, 'category_suggestion': category,
            'classification_confidence': confidence, 'audience_suggestion': audience,
            'extraction_status': extraction_status, 'extracted_chars': len(text),
            'processed_text_path': txt_rel, 'source_url': '', 'published_date': '',
            'last_verified': '', 'status': 'review' if flags else 'needs_source_verification',
            'review_flags': ';'.join(dict.fromkeys(flags)), 'processed_on': date.today().isoformat()
        })
    fieldnames = list(rows[0].keys()) if rows else ['document_id','filename','file_type','size_bytes','sha256','category_suggestion','classification_confidence','audience_suggestion','extraction_status','extracted_chars','processed_text_path','source_url','published_date','last_verified','status','review_flags','processed_on']
    with report.open('w', newline='', encoding='utf-8-sig') as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader(); writer.writerows(rows)
    print(f'Files scanned: {len(files)}')
    print(f'CSV report: {report}')
    print(f'Extracted text folder: {out}')
    print('IMPORTANT: category suggestions and duplicate flags require review; this script does not verify legal/current status.')
    print('Files needing review:', sum(1 for r in rows if r['review_flags']))

if __name__ == '__main__': main()
